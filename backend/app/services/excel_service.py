import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter


def get_output_dir():
    base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    output_dir = os.path.join(base_dir, "uploads", "excel")
    os.makedirs(output_dir, exist_ok=True)
    return output_dir


def generate_invoice_excel(invoice_data, output_dir=None):
    """Generate professional Excel invoice report"""
    if output_dir is None:
        output_dir = get_output_dir()
    
    os.makedirs(output_dir, exist_ok=True)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice Details"
    
    header_font = Font(name='Arial', size=14, bold=True, color="FFFFFF")
    sub_header_font = Font(name='Arial', size=11, bold=True)
    normal_font = Font(name='Arial', size=10)
    title_font = Font(name='Arial', size=18, bold=True, color="1F4E79")
    
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    light_fill = PatternFill(start_color="D6E3F8", end_color="D6E3F8", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    thick_border = Border(
        left=Side(style='medium'),
        right=Side(style='medium'),
        top=Side(style='medium'),
        bottom=Side(style='medium')
    )
    
    def set_cell(row, col, value, font=None, fill=None, alignment=None, border=None):
        cell = ws.cell(row=row, column=col, value=value)
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        if alignment:
            cell.alignment = alignment
        if border:
            cell.border = border
        return cell
    
    ws.merge_cells('A1:G1')
    set_cell(1, 1, "INVOICE", title_font, white_fill, Alignment(horizontal='center', vertical='center'), thick_border)
    ws.row_dimensions[1].height = 40
    
    ws.merge_cells('A2:G2')
    set_cell(2, 1, "InvoiceIQ AI - Smart Invoice Management", 
             Font(name='Arial', size=10, italic=True, color="666666"), 
             white_fill, Alignment(horizontal='center'))
    
    ws.merge_cells('A4:D4')
    set_cell(4, 1, "VENDOR INFORMATION", header_font, header_fill, Alignment(horizontal='left'), thin_border)
    
    vendor_info = [
        ("Vendor Name:", invoice_data.get('vendor_name', 'N/A')),
        ("GST Number:", invoice_data.get('gst_number', 'N/A')),
        ("Email:", invoice_data.get('email', 'N/A')),
        ("Phone:", invoice_data.get('phone_number', 'N/A'))
    ]
    
    for idx, (label, value) in enumerate(vendor_info, start=5):
        ws.merge_cells(f'B{idx}:D{idx}')
        set_cell(idx, 1, label, sub_header_font, light_fill, Alignment(horizontal='left'), thin_border)
        set_cell(idx, 2, value, normal_font, white_fill, Alignment(horizontal='left'), thin_border)
    
    ws.merge_cells('E4:G4')
    set_cell(4, 5, "INVOICE DETAILS", header_font, header_fill, Alignment(horizontal='left'), thin_border)
    
    invoice_details = [
        ("Invoice Number:", invoice_data.get('invoice_number', 'N/A')),
        ("Invoice Date:", invoice_data.get('invoice_date', 'N/A')),
        ("Total Amount:", f"₹ {invoice_data.get('total_amount', 0):,.2f}" if isinstance(invoice_data.get('total_amount'), (int, float)) else invoice_data.get('total_amount', 'N/A')),
        ("Generated:", datetime.now().strftime("%Y-%m-%d %H:%M"))
    ]
    
    for idx, (label, value) in enumerate(invoice_details, start=5):
        set_cell(idx, 5, label, sub_header_font, light_fill, Alignment(horizontal='left'), thin_border)
        set_cell(idx, 6, value, normal_font, white_fill, Alignment(horizontal='left'), thin_border)
        ws.merge_cells(f'F{idx}:G{idx}')
    
    row = 10
    ws.merge_cells(f'A{row}:G{row}')
    set_cell(row, 1, "PRODUCT DETAILS", header_font, header_fill, Alignment(horizontal='left'), thin_border)
    
    row += 1
    headers = ["S.No", "Item Name", "Quantity", "Unit Price", "Amount"]
    for col, header in enumerate(headers, start=1):
        set_cell(row, col, header, sub_header_font, light_fill, Alignment(horizontal='center'), thin_border)
    ws.merge_cells(f'E{row}:G{row}')
    
    row += 1
    products = invoice_data.get('products', [])
    if not products:
        products = [{"item_name": "No items", "quantity": "-", "amount": 0}]
    
    for idx, product in enumerate(products, start=1):
        amount = product.get('amount', 0)
        quantity = product.get('quantity', '-')
        
        if isinstance(quantity, str) and quantity.replace('.', '').replace('-', '').isdigit():
            quantity_val = float(quantity)
            amount_val = float(amount) if amount else 0
            unit_price = amount_val / quantity_val if quantity_val > 0 else 0
        else:
            amount_val = float(amount) if amount else 0
            unit_price = amount_val
        
        set_cell(row, 1, idx, normal_font, white_fill, Alignment(horizontal='center'), thin_border)
        set_cell(row, 2, product.get('item_name', 'N/A'), normal_font, white_fill, Alignment(horizontal='left'), thin_border)
        set_cell(row, 3, str(quantity), normal_font, white_fill, Alignment(horizontal='center'), thin_border)
        set_cell(row, 4, f"₹ {unit_price:,.2f}", normal_font, white_fill, Alignment(horizontal='right'), thin_border)
        set_cell(row, 5, f"₹ {amount_val:,.2f}", normal_font, white_fill, Alignment(horizontal='right'), thin_border)
        ws.merge_cells(f'E{row}:G{row}')
        row += 1
    
    row += 1
    total = invoice_data.get('total_amount', 0)
    if isinstance(total, str):
        try:
            total = float(total.replace(',', ''))
        except:
            total = 0
    
    set_cell(row, 1, "", normal_font, light_fill, Alignment(horizontal='center'), thin_border)
    set_cell(row, 2, "", normal_font, light_fill, Alignment(horizontal='left'), thin_border)
    ws.merge_cells(f'A{row}:C{row}')
    
    set_cell(row, 4, "GRAND TOTAL:", sub_header_font, header_fill, Alignment(horizontal='right'), thin_border)
    set_cell(row, 5, f"₹ {total:,.2f}", Font(name='Arial', size=12, bold=True, color="FFFFFF"), header_fill, Alignment(horizontal='right'), thick_border)
    ws.merge_cells(f'E{row}:G{row}')
    
    row += 2
    ws.merge_cells(f'A{row}:G{row}')
    set_cell(row, 1, "This invoice was generated by InvoiceIQ AI - Automated Invoice Management System",
             Font(name='Arial', size=9, italic=True, color="888888"), 
             white_fill, Alignment(horizontal='center'))
    
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15
    
    filename = f"invoice_{invoice_data.get('invoice_number', 'unknown').replace('/', '-')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(output_dir, filename)
    
    wb.save(filepath)
    
    return filepath